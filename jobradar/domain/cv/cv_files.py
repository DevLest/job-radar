"""CV files on disk: safe storage of uploads and local (free) text extraction."""

from __future__ import annotations

import re
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class CvFileStore:
    def __init__(self, directory: Path):
        self.directory = directory

    def save(self, filename: str, data: bytes) -> Path:
        self.directory.mkdir(parents=True, exist_ok=True)
        safe_name = re.sub(r"[^\w.\- ]+", "_", Path(filename).name).strip() or "cv.pdf"
        if Path(safe_name).suffix.lower() not in ALLOWED_EXTENSIONS:
            raise ValueError(f"Unsupported file type - upload one of {', '.join(sorted(ALLOWED_EXTENSIONS))}")
        path = self.directory / safe_name
        path.write_bytes(data)
        return path


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(path).pages).strip()
    if extension == ".docx":
        import docx
        document = docx.Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    return path.read_text(encoding="utf-8", errors="ignore").strip()
