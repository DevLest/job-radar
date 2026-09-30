"""CV upload + analysis. The CV text is extracted locally (free) and analyzed once;
the result can fill in your profile, and the text is reused for application drafts."""

from __future__ import annotations

import base64
import json
import re
from pathlib import Path

from . import llm
from .db import now
from .settings import CV_DIR

ALLOWED = {".pdf", ".docx", ".txt", ".md"}

SCHEMA = {
    "type": "object",
    "properties": {
        "full_name": {"type": "string"},
        "email": {"type": "string"},
        "phone": {"type": "string"},
        "headline": {"type": "string", "description": "Best-fit job title, e.g. 'Full-stack Laravel Developer'"},
        "years_experience": {"type": "number"},
        "seniority": {"type": "string", "enum": ["junior", "mid", "senior", "lead"]},
        "summary": {"type": "string", "description": "2-3 sentence professional summary in third person"},
        "based_in": {"type": "string"},
        "languages_spoken": {"type": "array", "items": {"type": "string"}},
        "primary_skills": {"type": "array", "items": {"type": "string"},
                           "description": "Core skills with real work evidence, lowercase, as job ads spell them (max 10)"},
        "secondary_skills": {"type": "array", "items": {"type": "string"}, "description": "Supporting skills (max 15)"},
        "learning_skills": {"type": "array", "items": {"type": "string"}, "description": "Mentioned but little evidence"},
        "title_keywords": {"type": "array", "items": {"type": "string"},
                           "description": "Lowercase words that appear in titles of jobs this person should see"},
        "strengths": {"type": "array", "items": {"type": "string"}},
        "gaps": {"type": "array", "items": {"type": "string"}, "description": "What employers will see as missing"},
        "cv_tips": {"type": "array", "items": {"type": "string"}, "description": "Concrete CV improvements (max 6)"},
    },
    "required": ["full_name", "email", "phone", "headline", "years_experience", "seniority", "summary", "based_in",
                 "languages_spoken", "primary_skills", "secondary_skills", "learning_skills", "title_keywords",
                 "strengths", "gaps", "cv_tips"],
    "additionalProperties": False,
}

SYSTEM = """You are an experienced tech recruiter. Analyze the candidate's CV and return a strict JSON profile.
Only claim skills the CV supports; put weakly-evidenced skills under learning_skills. Estimate years of
professional experience from dates. Be direct and specific in gaps and cv_tips."""


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages).strip()
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    return path.read_text(encoding="utf-8", errors="ignore").strip()


def save_upload(filename: str, data: bytes) -> Path:
    CV_DIR.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w.\- ]+", "_", Path(filename).name).strip() or "cv.pdf"
    if Path(safe).suffix.lower() not in ALLOWED:
        raise ValueError(f"Unsupported file type - upload one of {', '.join(sorted(ALLOWED))}")
    path = CV_DIR / safe
    path.write_bytes(data)
    return path


def analyze(db, profile: dict, path: Path) -> dict:
    text = extract_text(path)
    if len(text) >= 200:
        content = f"CV ({path.name}):\n\n{text[:40000]}"
    elif path.suffix.lower() == ".pdf":  # scanned PDF - let Claude read the pages directly
        content = [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                   "data": base64.b64encode(path.read_bytes()).decode()}},
                   {"type": "text", "text": "Analyze this CV."}]
    else:
        raise ValueError("Could not read any text from this file.")
    analysis = llm.ask(db, "cv", profile, SYSTEM, content, SCHEMA, max_tokens=8000, effort="medium")
    db.exec("UPDATE cvs SET active=0")
    db.exec("INSERT INTO cvs (filename, path, uploaded_at, text, analysis, active) VALUES (?,?,?,?,?,1)",
            (path.name, str(path), now(), text, json.dumps(analysis)))
    return analysis


def active(db):
    return db.one("SELECT * FROM cvs WHERE active=1 ORDER BY id DESC LIMIT 1")


def apply_to_profile(profile: dict, a: dict, replace_skills: bool = False) -> dict:
    cand = profile.setdefault("candidate", {})
    cand.update(title=a["headline"], years_experience=a["years_experience"], seniority=a["seniority"],
                summary=a["summary"])
    if a["full_name"]:
        cand["name"] = a["full_name"]
    if a["based_in"]:
        cand["based_in"] = a["based_in"]
    if a["languages_spoken"]:
        cand["languages_spoken"] = a["languages_spoken"]
    skills = profile.setdefault("skills", {})
    for key, src in (("primary", "primary_skills"), ("secondary", "secondary_skills"), ("learning", "learning_skills")):
        new = [s.lower() for s in a[src]]
        skills[key] = new if replace_skills else list(dict.fromkeys((skills.get(key) or []) + new))
    filters = profile.setdefault("filters", {})
    filters["title_include"] = list(dict.fromkeys((filters.get("title_include") or []) + a["title_keywords"]))
    return profile
